# -*- coding: utf-8 -*-
# K05 解析与质量分级 单元测试（**用解析工具链 venv** 运行）：
# 用法: python tests/test_corpus_parsers.py          （全部用例）
#       python tests/test_corpus_parsers.py lazy      （按用例函数名过滤单例，服务 venv 可跑该例）
# 覆盖：DOCX/XLSX/PPTX/PDF/CSV/JSON/TXT/MD 现场构造样本解析、旧 PPT 签名识别与受控转换、
#       ZIP 成员清单与体积上限、扫描 PDF 标 needs_ocr、图片/扫描页 OCR、质量分级与损坏文件、
#       惰性导入纪律（import 阶段零重依赖；缺依赖时 RuntimeError 提示工具链 venv）。
#
# OCR 与旧 PPT 转换用例**并入本文件**（不另建 test_corpus_ocr.py）：二者与被测单元同为
# corpus.parsers，共用 build_docx/build_pptx/build_text_pdf 等现场构造器与本文件的自定义
# 断言风格；惰性导入用例要求"用例最先运行且此刻无重库已加载"，分文件反而要重复该前置。
import hashlib
import os
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

RESULTS = []


def record(name, ok, detail=""):
    ok = bool(ok)
    RESULTS.append(ok)
    print(f"[{'PASS' if ok else 'FAIL'}] {name}" + (f" — {detail}" if detail and not ok else ""))


def tmp_dir(prefix):
    base = Path("D:/claudecode/scratch/sm-corpus-20261019")  # 只在 D 盘落临时产物（铁律 7）
    base.mkdir(parents=True, exist_ok=True)
    return Path(tempfile.mkdtemp(prefix=prefix, dir=str(base)))


# ── 样本构造（不依赖外部素材） ────────────────────────────────────────────────

def build_docx(path):
    import docx
    d = docx.Document()
    d.add_heading("动火作业检查表", level=1)
    d.add_paragraph("本表用于动火作业前的安全检查与签字确认，由作业负责人填写。")
    t = d.add_table(rows=3, cols=3)
    for i, h in enumerate(["序号", "检查项", "签字"]):
        t.rows[0].cells[i].text = h
    t.rows[1].cells[0].text = "1"
    t.rows[1].cells[1].text = "动火分析合格"
    t.rows[1].cells[2].text = ""
    t.rows[1].cells[1].merge(t.rows[1].cells[2])     # 横向合并，保留合并关系
    t.rows[2].cells[0].text = "2"
    t.rows[2].cells[1].text = "监护人到位"
    t.rows[2].cells[2].text = "张三"
    for i in range(12):
        d.add_paragraph(f"第{i + 1}条 动火作业票有效期与监护人职责说明，尾部不应被截断。")
    d.add_paragraph("尾部标记：动火作业记录保存三年。")   # 用于验证"完整正文不截断"
    d.add_paragraph("签字栏：作业人、监护人、审批人签字。")
    d.save(str(path))


def build_docx_tiny(path):
    import docx
    d = docx.Document()
    d.add_paragraph("短")     # 故意过短：正文 < 30 字符
    d.save(str(path))


def build_xlsx(path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "动火票"
    ws["A1"] = "序号"; ws["B1"] = "检查项(℃)"; ws["C1"] = "结论"
    for i in range(2, 12):
        ws[f"A{i}"] = i - 1
        ws[f"B{i}"] = f"动火分析第{i - 1}项合格"
        ws[f"C{i}"] = "是"
    ws["D14"] = "=SUM(A2:A11)"
    ws.merge_cells("A1:A2")
    ws2 = wb.create_sheet("受限空间")
    ws2["A1"] = "序号"; ws2["B1"] = "氧含量(%)"
    for i in range(2, 8):
        ws2[f"A{i}"] = i - 1
        ws2[f"B{i}"] = "19.5"
    wb.save(str(path))


def build_pptx(path):
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "动火作业票办理要点"
    slide.placeholders[1].text = "动火分析合格后方可作业，监护人不得离开现场。"
    slide.notes_slide.notes_text_frame.text = "讲师备注：强调动火票有效期与监护人职责。"
    prs.save(str(path))


def build_text_pdf(path, text):
    """现场构造最小文本 PDF（标准字体 + 内容流），pypdf 可提取文字。"""
    stream = f"BT /F1 14 Tf 40 150 Td ({text}) Tj ET".encode("latin-1")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 400 200] "
         b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"),
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode() + b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n"
            + str(xref) + "\n%%EOF\n").encode()
    path.write_bytes(bytes(out))


def build_blank_pdf(path):
    """无文字层 PDF（模拟扫描件，无 OCR 结果）。"""
    from pypdf import PdfWriter
    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    with open(path, "wb") as f:
        w.write(f)


def build_zip(path, members, big=False):
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, payload in members.items():
            z.writestr(name, payload)


def build_legacy_ppt(path):
    """OLE 复合文档签名文件（旧 .ppt 真实签名 D0CF11E0）。"""
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512)


def cjk_font_path():
    """本机中文字体路径（只读取系统字体，不落盘、不下载）；找不到返回 None。"""
    for candidate in ("C:/Windows/Fonts/msyh.ttc", "C:/Windows/Fonts/msyhbd.ttc",
                      "C:/Windows/Fonts/simhei.ttf", "C:/Windows/Fonts/simsun.ttc"):
        if Path(candidate).exists():
            return candidate
    return None


def draw_cjk_image(path, text, font_path, *, font_size=48, jpeg=False):
    """用 Pillow 现场画一张带中文文字的图片（不依赖任何外部素材）。返回 (宽, 高)。"""
    from PIL import Image, ImageDraw, ImageFont
    font = ImageFont.truetype(str(font_path), font_size)
    width = max(400, int(font_size * (len(text) + 2)))
    height = int(font_size * 3)
    image = Image.new("RGB", (width, height), "white")
    ImageDraw.Draw(image).text((font_size // 2, font_size), text, fill="black", font=font)
    if jpeg:
        image.save(path, "JPEG", quality=92)
    else:
        image.save(path)
    return width, height


def build_mixed_pdf(path, specs):
    """手写多页 PDF：每页为文本层（Helvetica 内容流）或无文字层的 JPEG 位图 XObject。

    specs 元素：{"text": "..."} 或 {"jpeg": b"...", "iw": N, "ih": M, "w": .., "h": ..}。
    """
    total = sum(2 + (1 if spec.get("jpeg") else 0) for spec in specs)
    font_no = 3 + total                      # 字体对象排在所有页对象之后
    objs = {1: b"<< /Type /Catalog /Pages 2 0 R >>"}
    kids = []
    n = 3
    for spec in specs:
        page_no, content_no = n, n + 1
        n += 2
        img_no = None
        if spec.get("jpeg"):
            img_no = n
            n += 1
        w = int(spec.get("w", 400))
        h = int(spec.get("h", 200))
        if img_no is None:
            stream = ("BT /F1 14 Tf 40 150 Td (%s) Tj ET" % spec["text"]).encode("latin-1")
            objs[content_no] = (b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n"
                                + stream + b"\nendstream")
            res = b"/Font << /F1 " + str(font_no).encode() + b" 0 R >>"
        else:
            data = spec["jpeg"]
            content = ("q %d 0 0 %d 0 0 cm /Im0 Do Q" % (w, h)).encode("latin-1")
            objs[content_no] = (b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n"
                                + content + b"\nendstream")
            objs[img_no] = (b"<< /Type /XObject /Subtype /Image /Width " + str(spec["iw"]).encode()
                            + b" /Height " + str(spec["ih"]).encode()
                            + b" /ColorSpace /DeviceRGB /BitsPerComponent 8 /Filter /DCTDecode"
                            + b" /Length " + str(len(data)).encode() + b" >>\nstream\n"
                            + data + b"\nendstream")
            res = b"/XObject << /Im0 " + str(img_no).encode() + b" 0 R >>"
        objs[page_no] = (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 " + str(w).encode() + b" "
                         + str(h).encode() + b"] /Resources << " + res + b" >> /Contents "
                         + str(content_no).encode() + b" 0 R >>")
        kids.append(page_no)
    objs[2] = (b"<< /Type /Pages /Kids [" + b" ".join(str(k).encode() + b" 0 R" for k in kids)
               + b"] /Count " + str(len(kids)).encode() + b" >>")
    objs[font_no] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"

    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for i in range(1, max(objs) + 1):
        offsets[i] = len(out)
        out += f"{i} 0 obj\n".encode() + objs[i] + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {max(objs) + 1}\n".encode() + b"0000000000 65535 f \n"
    for i in range(1, max(objs) + 1):
        out += f"{offsets[i]:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {max(objs) + 1} /Root 1 0 R >>\nstartxref\n"
            + str(xref) + "\n%%EOF\n").encode()
    path.write_bytes(bytes(out))


# ── 用例 ──────────────────────────────────────────────────────────────────────

def test_lazy_import_discipline():
    """先跑：import 阶段不得拉起重库；缺依赖时必须 RuntimeError 提示工具链 venv。"""
    heavy = [m for m in ("docx", "openpyxl", "pptx", "pypdf") if m in sys.modules]
    record("import 前无重解析库已加载（用例最先运行）", not heavy, f"heavy={heavy}")
    import corpus.parsers as P  # noqa: N812
    record("import corpus.parsers 零重依赖",
           all(m not in sys.modules for m in ("docx", "openpyxl", "pptx", "pypdf")))

    tmp = tmp_dir("lazy_")
    try:
        import importlib.util
        has_docx = importlib.util.find_spec("docx") is not None
        docx_path = tmp / "sample.docx"
        if has_docx:
            build_docx(docx_path)          # 工具链 venv：真实样本 + 屏蔽导入模拟缺依赖
        else:
            docx_path.write_bytes(b"\x99" * 64)   # 服务 venv：天然缺依赖
        txt_path = tmp / "note.txt"
        txt_path.write_text("动火作业票有效期与监护人职责说明，本文件解析不依赖任何解析工具链重库。",
                            encoding="utf-8")

        saved = {k: v for k, v in list(sys.modules.items())
                 if k == "docx" or k.startswith("docx.")}

        class _BlockDocx:
            def find_spec(self, name, path=None, target=None):
                if name == "docx":
                    raise ImportError("blocked for lazy-import test")
                return None

        blocker = _BlockDocx()
        if has_docx:
            sys.meta_path.insert(0, blocker)
            for k in list(saved):
                sys.modules.pop(k, None)
        try:
            try:
                P.parse_file(str(docx_path))
                msg = ""
                raised = False
            except RuntimeError as ex:
                raised = True
                msg = str(ex)
            except Exception as ex:  # noqa: BLE001
                raised = False
                msg = f"{type(ex).__name__}: {ex}"
            record("缺 docx 依赖时抛 RuntimeError 且提示工具链 venv",
                   raised and "工具链 venv" in msg, f"msg={msg}")
            txt_res = P.parse_file(str(txt_path))
            record("缺重库时文本类解析仍可用",
                   txt_res.status == "parsed" and txt_res.quality == "ok",
                   f"status={txt_res.status}")
        finally:
            if has_docx and blocker in sys.meta_path:
                sys.meta_path.remove(blocker)
            sys.modules.update(saved)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_docx_parse():
    from corpus.parsers import parse_file
    tmp = tmp_dir("docx_")
    try:
        p = tmp / "sample.docx"
        build_docx(p)
        res = parse_file(str(p))
        raw = p.read_bytes()
        text = "\n".join(b.content for b in res.blocks)
        table = next((b for b in res.blocks if b.kind == "table"), None)
        record("DOCX 解析成功且质量为 ok", res.status == "parsed" and res.quality == "ok",
               f"status={res.status} quality={res.quality} warnings={res.warnings}")
        record("DOCX 记录解析工具与版本",
               res.parser_name == "python-docx" and bool(res.parser_version))
        record("DOCX 保留标题/段落/表格三类块",
               any(b.kind == "heading" for b in res.blocks)
               and any(b.kind == "text" for b in res.blocks) and table is not None)
        record("DOCX 表格保留合并关系与签字栏",
               table is not None and bool(table.detail.get("merged_cells"))
               and table.detail.get("signature_block") is True,
               f"detail={table.detail if table else None}")
        record("DOCX 完整正文不截断（尾部段落可见）",
               "尾部标记：动火作业记录保存三年。" in text and "第12条" in text)
        record("DOCX 记录字节与 sha256",
               res.bytes == len(raw) and res.sha256 == hashlib.sha256(raw).hexdigest())
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_xlsx_parse():
    from corpus.parsers import parse_file
    tmp = tmp_dir("xlsx_")
    try:
        p = tmp / "sample.xlsx"
        build_xlsx(p)
        res = parse_file(str(p))
        sheet_blocks = [b for b in res.blocks if b.kind == "table"]
        joined = "\n".join(b.content for b in res.blocks)
        record("XLSX 解析成功且质量为 ok", res.status == "parsed" and res.quality == "ok",
               f"status={res.status} warnings={res.warnings}")
        record("XLSX 多 sheet 计数与定位",
               res.pages_or_sheets == 2 and {b.locator.get("sheet") for b in sheet_blocks} ==
               {"动火票", "受限空间"}, f"pages={res.pages_or_sheets}")
        record("XLSX 保留公式文本", any("=SUM(A2:A11)" in (b.detail or {}).get("formulas", [])
                                       for b in sheet_blocks))
        record("XLSX 保留合并单元格关系",
               any("A1:A2" in (b.detail or {}).get("merged_cells", []) for b in sheet_blocks))
        record("XLSX 识别表头与单位",
               any((b.detail or {}).get("units") for b in sheet_blocks)
               and any("序号" in (b.detail or {}).get("headers", []) for b in sheet_blocks))
        record("XLSX 行列值完整入正文", "氧含量(%)" in joined and "19.5" in joined)
        record("XLSX 使用 read_only 并记录工具版本",
               "read_only" in res.notes and res.parser_name == "openpyxl")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_pptx_parse():
    from corpus.parsers import parse_file
    tmp = tmp_dir("pptx_")
    try:
        p = tmp / "sample.pptx"
        build_pptx(p)
        res = parse_file(str(p))
        notes = [b for b in res.blocks if b.kind == "notes"]
        text = "\n".join(b.content for b in res.blocks)
        record("PPTX 解析成功且质量为 ok", res.status == "parsed" and res.quality == "ok",
               f"status={res.status} warnings={res.warnings}")
        record("PPTX 记录幻灯片页数与页码定位",
               res.pages_or_sheets == 1 and any(b.locator.get("slide") == 1 for b in res.blocks))
        record("PPTX 页内文字入正文", "动火分析合格后方可作业" in text)
        record("PPTX 保留备注/注释",
               notes and "讲师备注" in notes[0].content
               and notes[0].locator.get("slide") == 1)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_pdf_text_and_scan():
    from corpus.parsers import parse_file
    tmp = tmp_dir("pdf_")
    try:
        text_pdf = tmp / "text.pdf"
        build_text_pdf(text_pdf, "Hot work permit GB/T 30871-2022 requires a fire watch gas test.")
        r1 = parse_file(str(text_pdf))
        record("PDF 逐页提取文字成功",
               r1.status == "parsed" and "Hot work permit" in r1.blocks[0].content,
               f"status={r1.status} warnings={r1.warnings}")
        record("PDF 记录页码定位与页数",
               r1.pages_or_sheets == 1 and r1.blocks[0].locator == {"page": 1})

        scan_pdf = tmp / "scan.pdf"
        build_blank_pdf(scan_pdf)
        r2 = parse_file(str(scan_pdf))
        record("扫描型 PDF 页标 needs_ocr",
               r2.blocks and r2.blocks[0].detail.get("needs_ocr") is True
               and r2.blocks[0].locator.get("page") == 1, f"blocks={r2.blocks}")
        record("扫描型 PDF 整体不是 parsed",
               r2.status != "parsed" and r2.quality == "low_quality"
               and not r2.is_reliable_evidence(), f"status={r2.status} quality={r2.quality}")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_text_like_parse():
    from corpus.parsers import parse_file
    tmp = tmp_dir("text_")
    try:
        txt = tmp / "note.txt"
        txt.write_text("动火作业票有效期：特殊动火不超过8小时，一级动火不超过24小时。\n\n"
                       "监护人不得离开作业现场，发现异常立即停止作业并撤离。", encoding="utf-8")
        r1 = parse_file(str(txt))
        record("TXT 解析成功（段落定位）",
               r1.status == "parsed" and r1.kind == "txt"
               and any(b.locator.get("para") for b in r1.blocks),
               f"status={r1.status} kind={r1.kind}")

        md = tmp / "doc.md"
        md.write_text("# 动火作业安全管理规定\n\n依据 GB/T 30871-2022 执行动火票审批。\n"
                      "## 监护人职责\n\n监护人全程在场，不得离岗。", encoding="utf-8")
        r2 = parse_file(str(md))
        record("MD 解析成功并保留标题块",
               r2.status == "parsed" and any(b.kind == "heading" for b in r2.blocks))
        record("MD 正文完整", any("监护人全程在场" in b.content for b in r2.blocks))

        csv_path = tmp / "list.csv"
        csv_path.write_text("序号,检查项,结论\n1,动火分析合格,是\n2,监护人到位,是\n",
                            encoding="utf-8")
        r3 = parse_file(str(csv_path))
        table = r3.blocks[0] if r3.blocks else None
        record("CSV 解析为表格块并保留表头",
               r3.status == "parsed" and table and table.kind == "table"
               and table.detail.get("headers") == ["序号", "检查项", "结论"],
               f"detail={table.detail if table else None}")

        js = tmp / "items.json"
        js.write_text('[{"title": "动火作业检查表", "content": "动火前检查分析合格并确认监护人到位。"},'
                      ' {"title": "受限空间作业票", "content": "先通风、再检测、后作业。"}]',
                      encoding="utf-8")
        r4 = parse_file(str(js))
        record("JSON 数组解析为多条块并保留 title",
               r4.status == "parsed" and len(r4.blocks) == 2
               and r4.blocks[0].detail.get("title") == "动火作业检查表",
               f"blocks={len(r4.blocks)}")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_legacy_ppt_pending_conversion():
    from corpus.parsers import parse_file
    tmp = tmp_dir("ppt_")
    try:
        p = tmp / "old.ppt"
        build_legacy_ppt(p)
        res = parse_file(str(p))
        raw = p.read_bytes()
        record("旧 PPT 按真实签名识别（非当失败）",
               res.kind == "ppt_legacy" and res.status == "pending_conversion",
               f"kind={res.kind} status={res.status}")
        record("旧 PPT 保留原件 sha256/字节",
               res.sha256 == hashlib.sha256(raw).hexdigest() and res.bytes == len(raw))
        record("旧 PPT 标注待受控转换（含哈希/工具版本/页数对照）",
               "待受控转换" in res.notes and "哈希" in res.notes and res.parser_name == "ole_signature")
        record("旧 PPT 不冒出正文块", res.blocks == [])
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_zip_members_and_limit():
    from corpus.parsers import parse_file
    tmp = tmp_dir("zip_")
    try:
        p = tmp / "pack.zip"
        build_zip(p, {"动火作业表.txt": "动火", "macro.docm": "x" * 100})
        res = parse_file(str(p))
        block = res.blocks[0] if res.blocks else None
        members = {n for n, _s in (block.detail or {}).get("members", [])} if block else set()
        record("ZIP 解析成功（成员清单）",
               res.status == "parsed" and block and block.kind == "manifest",
               f"status={res.status}")
        record("ZIP 列出成员与解压后大小",
               members == {"动火作业表.txt", "macro.docm"}
               and (block.detail or {}).get("total_uncompressed") ==
               sum(s for _n, s in (block.detail or {}).get("members", [])),
               f"detail={block.detail if block else None}")
        record("ZIP 声明未解压/未执行宏",
               "未解压" in res.notes and any("宏" in w for w in res.warnings))

        big = tmp / "big.zip"
        build_zip(big, {"a.bin": "z" * 4096})
        os.environ["SAFETYMIND_ZIP_MAX_BYTES"] = "16"
        try:
            r2 = parse_file(str(big))
        finally:
            os.environ.pop("SAFETYMIND_ZIP_MAX_BYTES", None)
        record("ZIP 超受控体积上限报 unparseable 并说明",
               r2.status == "unparseable" and any("上限" in e for e in r2.errors),
               f"status={r2.status} errors={r2.errors}")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_quality_grades_and_corrupt():
    from corpus.parsers import parse_file
    tmp = tmp_dir("qual_")
    try:
        tiny = tmp / "tiny.docx"
        build_docx_tiny(tiny)
        r1 = parse_file(str(tiny))
        record("正文过短 → low_quality 且非可靠证据",
               r1.status == "low_quality" and r1.quality == "low_quality"
               and not r1.is_reliable_evidence(), f"status={r1.status}")

        garbled = tmp / "garbled.txt"
        garbled.write_bytes(b"\x01\x02" * 200)      # 控制字符占比高 → 乱码
        r2 = parse_file(str(garbled))
        record("乱码率高 → low_quality",
               r2.status == "low_quality" and any("乱码" in w for w in r2.warnings),
               f"status={r2.status} warnings={r2.warnings}")

        corrupt = tmp / "corrupt.docx"
        corrupt.write_bytes(b"\x99" * 256)          # 损坏 Office 文件（无 OOXML 标记）
        r3 = parse_file(str(corrupt))
        record("损坏文件 → unparseable（不抛未捕获异常）",
               r3.status == "unparseable" and r3.quality == "unparseable" and bool(r3.errors),
               f"status={r3.status} errors={r3.errors}")
        record("损坏文件仍保留哈希/字节",
               r3.sha256 == hashlib.sha256(b"\x99" * 256).hexdigest() and r3.bytes == 256)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_unsupported_and_image():
    from corpus.parsers import parse_file
    tmp = tmp_dir("misc_")
    try:
        other = tmp / "note.rtf"
        other.write_text("{\\rtf1 动火作业}", encoding="utf-8")
        r1 = parse_file(str(other))
        record("未支持格式 → unsupported（保留原件信息）",
               r1.status == "unsupported" and r1.quality == "unsupported"
               and bool(r1.sha256), f"status={r1.status} kind={r1.kind}")

        img = tmp / "scan.png"
        img.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 64)
        r2 = parse_file(str(img))
        record("图片 → image_scan 且标 needs_ocr（低质量）",
               r2.kind == "image_scan" and r2.status == "low_quality"
               and r2.blocks[0].detail.get("needs_ocr") is True,
               f"status={r2.status} kind={r2.kind}")

        missing = tmp / "nope.docx"
        try:
            parse_file(str(missing))
            record("不存在的路径抛 FileNotFoundError", False, "未抛错")
        except FileNotFoundError:
            record("不存在的路径抛 FileNotFoundError", True)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


class CountingOcrEngine:
    """假 OCR 引擎：记录每次识别调用的图像尺寸，返回固定长文本。

    用于证明"OCR 只在确实无文字层处触发"——有文字层的页必须零调用。
    """
    engine_name = "fake-ocr"
    engine_version = "0.0-test"
    lang = "ch"

    def __init__(self):
        self.calls = []

    def recognize(self, image):
        self.calls.append(getattr(image, "size", None))
        return ["合成识别文本：动火作业票有效期不超过8小时，监护人不得离开作业现场并全程监护。"]


class _FakeProc:
    """假转换进程结果（returncode/stdout/stderr）。"""

    def __init__(self, returncode=0, stdout=b"", stderr=b""):
        self.returncode = returncode
        self.stdout = stdout
        self.stderr = stderr


def test_ocr_trigger_scope():
    """OCR 触发范围与降级语义：仅无文字层处触发；引擎缺失如实降级不崩溃。"""
    from corpus import parsers as P
    tmp = tmp_dir("ocrscope_")
    font = cjk_font_path()
    if font is None:
        record("OCR 用例需要本机 CJK 字体以现场构造中文样本", False, "未找到 msyh/simhei/simsun")
        shutil.rmtree(tmp, ignore_errors=True)
        return
    engine = CountingOcrEngine()
    saved_factory = P._create_ocr_engine
    P._create_ocr_engine = lambda: (engine, "")
    P.reset_ocr_engine_cache()
    try:
        jpeg = tmp / "scan.jpg"
        iw, ih = draw_cjk_image(jpeg, "动火作业票有效期不超过8小时", font, jpeg=True)
        mixed = tmp / "mixed.pdf"
        build_mixed_pdf(mixed, [
            {"text": "Hot work permit GB/T 30871-2022 fire watch gas test requirement."},
            {"jpeg": jpeg.read_bytes(), "iw": iw, "ih": ih, "w": iw // 2, "h": ih // 2},
        ])
        res = P.parse_file(str(mixed))
        first, second = res.blocks[0], res.blocks[-1]
        record("有文字层的 PDF 页不被 OCR（引擎零调用、块无 ocr 标记）",
               first.kind == "text" and "Hot work permit" in first.content
               and not (first.detail or {}).get("ocr") and len(engine.calls) == 1,
               f"calls={len(engine.calls)} first={first.locator}")
        record("无文字层页触发 OCR 并产出 kind=text 文字块（保留页码定位）",
               second.kind == "text" and second.locator == {"page": 2}
               and "合成识别文本" in second.content, f"locator={second.locator}")
        ocr = (second.detail or {}).get("ocr") or {}
        record("OCR 结果块记录引擎名/版本/语言",
               {"engine", "version", "lang"} <= set(ocr)
               and ocr.get("engine") == "fake-ocr" and ocr.get("version") == "0.0-test"
               and ocr.get("lang") == "ch", f"ocr={ocr}")
        record("扫描页经 OCR 补齐后整体可判为 parsed/ok",
               res.status == "parsed" and res.quality == "ok" and res.pages_or_sheets == 2,
               f"status={res.status} quality={res.quality} warnings={res.warnings}")

        # 引擎缺失（注入返回 None 的工厂）：如实降级，不崩溃
        P._create_ocr_engine = lambda: (None, "模拟缺引擎：rapidocr 未安装且系统无 tesseract")
        P.reset_ocr_engine_cache()
        missing = P.parse_file(str(mixed))
        scan_blocks = [b for b in missing.blocks if (b.detail or {}).get("needs_ocr")]
        record("OCR 引擎缺失时扫描页如实降级（保留 needs_ocr/原因，不崩溃）",
               missing.status != "parsed" and missing.quality == "low_quality"
               and not missing.is_reliable_evidence() and bool(scan_blocks)
               and "模拟缺引擎" in (scan_blocks[0].detail or {}).get("ocr_unavailable", "")
               and any("OCR 引擎不可用" in w for w in missing.warnings),
               f"status={missing.status} warnings={missing.warnings}")

        img = tmp / "note.png"
        draw_cjk_image(img, "动火作业票有效期", font)
        img_missing = P.parse_file(str(img))
        record("引擎缺失时 image_scan 如实降级为待 OCR（不冒充已解析）",
               img_missing.kind == "image_scan" and img_missing.status == "low_quality"
               and img_missing.blocks[0].detail.get("needs_ocr") is True,
               f"status={img_missing.status}")
    finally:
        P._create_ocr_engine = saved_factory
        P.reset_ocr_engine_cache()
        shutil.rmtree(tmp, ignore_errors=True)


def test_ocr_real_engine_sample():
    """真实 OCR 引擎：图片路径与"仅含位图、无文字层"的 PDF 页都要产出可用文字块。"""
    from corpus import parsers as P
    tmp = tmp_dir("ocrreal_")
    try:
        engine, reason = P._resolve_ocr_engine()
        if engine is None:
            broken = tmp / "note.png"
            broken.write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 64)
            res = P.parse_file(str(broken))
            record("OCR 引擎不可用时图片解析如实降级（不崩溃、不冒充已解析）",
                   res.kind == "image_scan" and res.status == "low_quality"
                   and any("OCR 引擎不可用" in w for w in res.warnings), f"reason={reason}")
            return
        font = cjk_font_path()
        if font is None:
            record("OCR 真实识别用例需要本机 CJK 字体", False, "未找到 msyh/simhei/simsun")
            return
        # 文本须长于 MIN_TEXT_CHARS，才能验证"OCR 后仍可判为 ok"（短文本按既有规则本应 low_quality）
        expected = "动火作业票有效期不超过8小时，监护人不得离开作业现场，动火分析合格后方可作业"

        png = tmp / "scan.png"
        draw_cjk_image(png, expected, font, font_size=40)
        r1 = P.parse_file(str(png))
        t1 = r1.blocks[0].content if r1.blocks else ""
        print(f"      [OCR 样例·图片] {t1!r}（{engine.engine_name} {engine.engine_version}）")
        record("图片 OCR 真实产出文字块（kind=text / page 定位 / 引擎记录）",
               r1.status == "parsed" and r1.quality == "ok" and r1.blocks[0].kind == "text"
               and r1.blocks[0].locator == {"page": 1}
               and (r1.blocks[0].detail or {}).get("ocr", {}).get("engine") == engine.engine_name,
               f"status={r1.status} quality={r1.quality} text={t1!r}")
        record("图片 OCR 识别内容与预期文本高度重合",
               sum(1 for ch in expected if ch in t1) >= len(expected) * 0.6,
               f"text={t1!r}")

        jpeg = tmp / "page.jpg"
        iw, ih = draw_cjk_image(jpeg, expected, font, font_size=40, jpeg=True)
        pdf = tmp / "scan.pdf"
        build_mixed_pdf(pdf, [{"jpeg": jpeg.read_bytes(), "iw": iw, "ih": ih,
                               "w": iw // 2, "h": ih // 2}])
        r2 = P.parse_file(str(pdf), staging_dir=str(tmp))
        t2 = r2.blocks[0].content if r2.blocks else ""
        print(f"      [OCR 样例·图片型 PDF] {t2!r}")
        record("无文字层 PDF 页被 OCR 且产出文字块（保留页码、不再标 needs_ocr）",
               r2.status == "parsed" and r2.quality == "ok" and r2.blocks[0].kind == "text"
               and r2.blocks[0].locator == {"page": 1}
               and "engine" in (r2.blocks[0].detail or {}).get("ocr", {})
               and not (r2.blocks[0].detail or {}).get("needs_ocr")
               and r2.is_reliable_evidence(),
               f"status={r2.status} quality={r2.quality} text={t2!r}")
        record("图片型 PDF 的 OCR 识别内容与预期文本高度重合",
               sum(1 for ch in expected if ch in t2) >= len(expected) * 0.6, f"text={t2!r}")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _fake_converter(mode, src, good_pptx):
    """构造假转换器与调用记录：mode ∈ {ok, fail, timeout, empty}。"""
    from corpus import parsers as P
    calls = []

    def runner(cmd, timeout):
        calls.append(list(cmd))
        if mode == "timeout":
            raise subprocess.TimeoutExpired(cmd, timeout)
        if mode == "fail":
            return _FakeProc(1, stderr=b"source file could not be loaded")
        if mode == "empty":
            return _FakeProc(0)
        fmt = cmd[cmd.index("--convert-to") + 1]
        out_dir = Path(cmd[cmd.index("--outdir") + 1])
        out_dir.mkdir(parents=True, exist_ok=True)
        if fmt == "pptx":
            shutil.copyfile(good_pptx, out_dir / f"{src.stem}.pptx")
        return _FakeProc(0)

    converter = P.PptConverter(executable="soffice-fake", probe=lambda: "soffice-fake",
                               runner=runner, timeout_s=7,
                               page_probe=lambda s, d: (2, "假页数探测（源 2 页）"))
    return converter, calls


def test_ppt_converter_fake_paths():
    """假转换器覆盖成功/失败/超时/无产物四条路径；任何异常都不得静默当成功。"""
    from corpus import parsers as P
    tmp = tmp_dir("pptconv_")
    saved_factory = P._ppt_converter_factory
    try:
        src = tmp / "old.ppt"
        build_legacy_ppt(src)
        src_sha = hashlib.sha256(src.read_bytes()).hexdigest()
        src_bytes = len(src.read_bytes())
        good_pptx = tmp / "good.pptx"
        build_pptx(good_pptx)

        converter, calls = _fake_converter("ok", src, good_pptx)
        P._ppt_converter_factory = lambda: converter
        res = P.parse_file(str(src), staging_dir=str(tmp))
        conv = res.conversion or {}
        text = "\n".join(b.content for b in res.blocks)
        record("假转换器成功 → 继续解析出页数与文字块",
               res.status == "parsed" and res.kind == "ppt_legacy" and res.quality == "ok"
               and res.pages_or_sheets == 1 and "动火分析合格后方可作业" in text,
               f"status={res.status} pages={res.pages_or_sheets} text={text[:60]!r}")
        record("成功路径保留原件 sha256/字节，转换只记录输出路径与输出哈希",
               res.sha256 == src_sha and res.bytes == src_bytes
               and os.path.exists(conv.get("output_path") or "")
               and conv.get("output_sha256") != src_sha,
               f"sha={res.sha256[:12]} out={conv.get('output_path')}")
        record("成功路径记录工具名/版本字段、页数对照与命令",
               conv.get("ok") is True and conv.get("tool_name") == "libreoffice-soffice"
               and "tool_version" in conv and conv.get("source_pages") == 2
               and conv.get("output_pages") == 1 and conv.get("exit_code") == 0
               and isinstance(conv.get("command"), list) and bool(conv.get("command")),
               f"conv={conv}")
        record("成功路径 notes 明示页数对照（源 2 → 转换后 1）",
               "页数对照=源 2 → 转换后 1" in res.notes, f"notes={res.notes[:140]}")
        record("成功路径确实执行了输出到 .pptx 的受控命令",
               any("--convert-to" in c and "pptx" in c and "--outdir" in c for c in calls),
               f"calls={len(calls)}")

        converter, _calls = _fake_converter("fail", src, good_pptx)
        P._ppt_converter_factory = lambda: converter
        res_fail = P.parse_file(str(src), staging_dir=str(tmp))
        record("假转换器失败 → 保持 pending_conversion 且如实记录原因",
               res_fail.status == "pending_conversion" and res_fail.quality == "pending"
               and res_fail.blocks == [] and not res_fail.is_reliable_evidence()
               and any("退出码 1" in e for e in res_fail.errors)
               and "退出码 1" in (res_fail.conversion or {}).get("reason", "")
               and res_fail.sha256 == src_sha,
               f"status={res_fail.status} errors={res_fail.errors}")
        record("失败路径 notes 保留待受控转换口径与哈希说明",
               "待受控转换" in res_fail.notes and "哈希" in res_fail.notes)

        converter, _calls = _fake_converter("timeout", src, good_pptx)
        P._ppt_converter_factory = lambda: converter
        res_to = P.parse_file(str(src), staging_dir=str(tmp))
        record("假转换器超时 → 保持 pending_conversion 且记录超时",
               res_to.status == "pending_conversion" and res_to.blocks == []
               and "超时" in (res_to.conversion or {}).get("reason", "")
               and any("超时" in e for e in res_to.errors),
               f"reason={(res_to.conversion or {}).get('reason')}")

        converter, _calls = _fake_converter("empty", src, good_pptx)
        P._ppt_converter_factory = lambda: converter
        res_empty = P.parse_file(str(src), staging_dir=str(tmp))
        record("退出码 0 但无产物 → 绝不静默当成功",
               res_empty.status == "pending_conversion" and res_empty.blocks == []
               and "未产出" in (res_empty.conversion or {}).get("reason", ""),
               f"reason={(res_empty.conversion or {}).get('reason')}")

        # 无 LibreOffice（探测返回 None）：如实登记未就绪，不下载安装
        no_tool = P.PptConverter(probe=lambda: None, runner=lambda cmd, t: _FakeProc(0))
        P._ppt_converter_factory = lambda: no_tool
        res_none = P.parse_file(str(src), staging_dir=str(tmp))
        record("无转换工具时如实登记待受控转换（缺 LibreOffice，不下载安装）",
               res_none.status == "pending_conversion" and not no_tool.available()
               and "LibreOffice" in res_none.notes and "待受控转换" in res_none.notes
               and (res_none.conversion or {}).get("ok") is False
               and res_none.sha256 == src_sha,
               f"status={res_none.status} notes={res_none.notes[:100]}")
    finally:
        P._ppt_converter_factory = saved_factory
        shutil.rmtree(tmp, ignore_errors=True)


def test_ppt_converter_real_environment():
    """真实环境如实登记 + 真实 .ppt 端到端转换（有 LibreOffice 则实测，无则如实登记未就绪）。"""
    from corpus import parsers as P
    tmp = tmp_dir("pptreal_")
    try:
        exe = P._probe_soffice()
        converter = P.PptConverter()
        record("实测探测结果如实（有则记录路径，无则登记未就绪；不下载不硬装）",
               (exe is None and not converter.available())
               or (bool(exe) and converter.available()), f"exe={exe!r}")

        legacy = tmp / "legacy.ppt"
        build_legacy_ppt(legacy)
        if not exe:
            res = P.parse_file(str(legacy), staging_dir=str(tmp))
            record("无 LibreOffice：缺工具即 pending + 明确原因（可一键转换的说明在 notes）",
                   res.status == "pending_conversion" and "LibreOffice" in res.notes
                   and (res.conversion or {}).get("ok") is False,
                   f"status={res.status} notes={res.notes[:100]}")
            return
        print(f"      [LibreOffice 实测] exe={exe} 版本={converter.tool_version()!r}")
        record("探测到转换工具时能读到工具版本", bool(converter.tool_version()),
               f"tool_version={converter.tool_version()!r}")

        # 用真实 soffice 把 pptx 转成真实 .ppt（OLE）作为输入样本
        pptx_src = tmp / "src.pptx"
        build_pptx(pptx_src)
        ppt_dir = tmp / "from_pptx"
        ppt_dir.mkdir(parents=True, exist_ok=True)
        ppt_path = ppt_dir / "src.ppt"
        proc = subprocess.run(
            [exe, "-env:UserInstallation=file:///D:/SafetyMindData/lo_profile",
             "--headless", "--norestore", "--convert-to", "ppt:MS PowerPoint 97",
             "--outdir", str(ppt_dir), str(pptx_src)],
            capture_output=True, timeout=180)
        built = ppt_path.exists() and ppt_path.stat().st_size > 0
        record("用真实 soffice 现场生成 .ppt 压力样本（pptx→ppt）", built,
               f"stderr={(proc.stderr or b'')[-200:]!r}")
        if not built:
            return
        head = ppt_path.read_bytes()[:8]
        record("现场样本是真 OLE 复合文档（旧 .ppt 签名 D0CF11E0）",
               head == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1", f"head={head.hex()}")

        res = P.parse_file(str(ppt_path), staging_dir=str(tmp))
        conv = res.conversion or {}
        text = "\n".join(b.content for b in res.blocks)
        record("真实 .ppt 受控转换成功 → 解析出页数与文字块",
               res.status == "parsed" and res.kind == "ppt_legacy"
               and res.pages_or_sheets == 1 and "动火分析合格后方可作业" in text,
               f"status={res.status} pages={res.pages_or_sheets} text={text[:60]!r}")
        record("真实转换记录工具名/版本/输出路径/输出哈希/原件哈希",
               conv.get("ok") is True and conv.get("tool_name") == "libreoffice-soffice"
               and bool(conv.get("tool_version"))
               and os.path.exists(conv.get("output_path") or "")
               and bool(conv.get("output_sha256"))
               and res.sha256 == hashlib.sha256(ppt_path.read_bytes()).hexdigest(),
               f"conv={conv}")
        record("真实转换给出页数对照（源 N → 转换后 M 均为实测值）",
               conv.get("source_pages") == 1 and conv.get("output_pages") == 1
               and "页数对照=源 1 → 转换后 1" in res.notes,
               f"source={conv.get('source_pages')} output={conv.get('output_pages')}")
        record("真实转换产物写在 D 盘受控 staging 目录内",
               Path(conv.get("output_path", "")).drive.upper() == "D:",
               f"output={conv.get('output_path')}")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


TESTS = [
    test_lazy_import_discipline,
    test_docx_parse,
    test_xlsx_parse,
    test_pptx_parse,
    test_pdf_text_and_scan,
    test_text_like_parse,
    test_legacy_ppt_pending_conversion,
    test_zip_members_and_limit,
    test_quality_grades_and_corrupt,
    test_unsupported_and_image,
    test_ocr_trigger_scope,
    test_ocr_real_engine_sample,
    test_ppt_converter_fake_paths,
    test_ppt_converter_real_environment,
]

if __name__ == "__main__":
    selected = sys.argv[1:]
    cases = [t for t in TESTS if not selected or any(s in t.__name__ for s in selected)]
    for fn in cases:
        try:
            fn()
        except Exception as ex:  # noqa: BLE001
            record(fn.__name__, False, f"异常: {type(ex).__name__}: {ex}")
    passed, total = sum(RESULTS), len(RESULTS)
    print(f"\n{passed}/{total} passed")
    sys.exit(0 if passed == total else 1)
