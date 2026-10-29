# -*- coding: utf-8 -*-
# 安全台账文件导出测试（2026-10-26 闭环第 3 轮新增能力）。
#
# 背景（用户指出的真实痛点）：安全员台账最终要交 Office/WPS 文件，
# 而原 safety_ledger_assistant 只会"告诉你要准备哪些文件"，不能产出文件本身。
#
# 覆盖：
#   1. 五类法定台账模板齐备且带法律依据
#   2. xlsx 真实产出：文件存在、可被 openpyxl 重新读回、表头与模板一致
#   3. 产出内容不是空壳：含标题/依据/表头，且有空白填写行（可打印手填）
#   4. docx 能力缺失时给出**可执行的补救建议**而不是模糊报错
#   5. 未知台账类型 / 未知格式 → 结构化失败（不抛未捕获异常）
#   6. 产物落 D 盘（铁律 7），不写 C 盘
# 用法: python tests/test_ledger_export.py
import os
import shutil
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

RESULTS = []


def record(name, ok, detail=""):
    RESULTS.append(ok)
    print(f"[{'PASS' if ok else 'FAIL'}] {name}" + (f" — {detail}" if detail and not ok else ""))


def main():
    from core.ledger_export import (LEDGER_TEMPLATES, capability_report, export,
                                     export_xlsx)

    # ── 1. 模板齐备性 ──────────────────────────────────────────────────
    required = {"hazard", "inspection", "training", "ticket", "emergency_drill"}
    record("五类法定台账模板齐备", required <= set(LEDGER_TEMPLATES),
           f"实际={list(LEDGER_TEMPLATES)}")
    no_basis = [k for k, v in LEDGER_TEMPLATES.items() if not v.get("legal_basis")]
    record("每类台账都有法律依据说明", not no_basis, f"缺依据={no_basis}")
    empty_cols = [k for k, v in LEDGER_TEMPLATES.items() if len(v.get("columns") or []) < 5]
    record("每类台账列定义不少于 5 列", not empty_cols, f"列不足={empty_cols}")

    # ── 2. xlsx 真实产出与可读回 ────────────────────────────────────────
    tmp = Path("D:/claudecode/scratch/sm-loop/test_ledger_out")
    if tmp.exists():
        shutil.rmtree(str(tmp), ignore_errors=True)
    tmp.mkdir(parents=True, exist_ok=True)

    try:
        import openpyxl
        p = export_xlsx("hazard", tmp, with_demo=True)
        record("xlsx 产出文件真实存在", p.exists() and p.stat().st_size > 1000,
               f"path={p} size={p.stat().st_size if p.exists() else 0}")
        wb = openpyxl.load_workbook(str(p))
        record("xlsx 可被 openpyxl 重新读回（不是损坏文件）", len(wb.sheetnames) >= 1,
               f"sheets={wb.sheetnames}")
        ws = wb[wb.sheetnames[0]]
        header = [ws.cell(row=3, column=c).value for c in range(1, ws.max_column + 1)]
        expect_cols = [c[0] for c in LEDGER_TEMPLATES["hazard"]["columns"]]
        record("xlsx 表头与模板列定义逐列一致", header == expect_cols,
               f"实际={header} 期望={expect_cols}")

        title = ws.cell(row=1, column=1).value
        basis = ws.cell(row=2, column=1).value
        record("xlsx 含台账标题行", title == LEDGER_TEMPLATES["hazard"]["name"], f"title={title}")
        record("xlsx 含法律依据行", bool(basis) and "安全生产法" in str(basis), f"basis={basis}")

        # 示例数据 + 空白填写行都要在
        demo_present = any(ws.cell(row=r, column=1).value == 1 for r in range(4, 10))
        blanks = sum(1 for r in range(4, ws.max_row + 1)
                     if ws.cell(row=r, column=1).value is None)
        record("xlsx 含示例数据行", demo_present)
        record("xlsx 留有空白填写行（可打印手填）", blanks >= 5, f"空白行={blanks}")
        record("xlsx 含填写说明页", "填写说明" in wb.sheetnames, f"sheets={wb.sheetnames}")
    except Exception as e:
        record("xlsx 产出链路", False, f"{type(e).__name__}: {e}")

    # ── 3. 无示例时也应有空白可填行（空模板场景）────────────────────────
    try:
        p2 = export_xlsx("inspection", tmp, with_demo=False)
        wb2 = openpyxl.load_workbook(str(p2))
        ws2 = wb2[wb2.sheetnames[0]]
        record("空模板也有表头与空白行", ws2.max_row >= 10 and ws2.cell(row=3, column=1).value == "序号",
               f"max_row={ws2.max_row}")
    except Exception as e:
        record("空模板生成", False, f"{type(e).__name__}: {e}")

    # ── 4. docx 能力缺失时给可执行补救建议 ──────────────────────────────
    caps = capability_report()
    r = export("inspection", "docx", tmp, False)
    if caps["docx"]:
        record("docx 可用时能真实产出", r.get("ok") is True and Path(r["path"]).exists(),
               f"r={r}")
    else:
        err = str(r.get("error") or "")
        record("docx 缺失时返回结构化失败", r.get("ok") is False, f"r={r}")
        record("docx 缺失的错误里给出可执行补救建议（装包路径/替代格式）",
               ("python-docx" in err) and ("xlsx" in err or "pip install" in err),
               f"err={err[:200]}")

    # ── 5. 未知台账/格式结构化失败 ──────────────────────────────────────
    r_bad_ledger = export("__no_such__", "xlsx", tmp)
    record("未知台账类型返回结构化失败并列出可用项",
           r_bad_ledger.get("ok") is False and len(r_bad_ledger.get("available") or []) >= 5,
           f"r={r_bad_ledger}")
    r_bad_fmt = export("hazard", "pdf", tmp)
    record("未知格式返回结构化失败", r_bad_fmt.get("ok") is False, f"r={r_bad_fmt}")

    # ── 6. 产物只落 D 盘（铁律 7）───────────────────────────────────────
    r_d = export("hazard", "xlsx", tmp, True)
    path_d = str(r_d.get("path") or "")
    record("产物落在 D 盘（不写 C 盘）", path_d.startswith("D:"), f"path={path_d}")

    # 清理
    shutil.rmtree(str(tmp), ignore_errors=True)

    failed = RESULTS.count(False)
    print(f"\n{len(RESULTS) - failed}/{len(RESULTS)} PASSED")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
