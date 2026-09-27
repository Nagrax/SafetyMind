# 将真实业务 xlsx（隐患排查依据大全 + 安环部工作清单）转换为知识库文档并入库
# 特性: 续行合并（子条目并入上一检查项）、前向填充类别、bge 语义嵌入入库
# 用法: python tools/seed_xlsx_kb.py [--dry-run]
import os, sys, json, time
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
os.chdir(os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
from dotenv import load_dotenv
load_dotenv(".env")

import openpyxl

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
XLSX1 = os.path.join(ROOT, "2026《隐患排查依据大全》查找版 (1).xlsx")
XLSX2 = os.path.join(ROOT, "安环部工作清单(1).xlsx")
SEED_JSON = os.path.join(ROOT, "data", "knowledge_seed", "xlsx_docs.json")

def clean(v):
    return str(v).replace("\xa0", " ").strip() if v is not None else ""

def convert_yiju(path):
    """隐患排查依据大全: 行 → 文档；续行（无大/小类）并入上一条"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    docs, cur = [], None
    for r in ws.iter_rows(min_row=3, values_only=True):
        if not any(v is not None for v in r):
            continue
        seq, big, mid, small, content, basis, level = (clean(r[i]) for i in range(7))
        content = content.strip()
        if big and not content:  # 纯表头行跳过
            continue
        if not any([big, mid, small]) and content and cur is not None:
            cur["content"] += f"\n{content}"          # 续行并入上一检查项
            continue
        if not content:
            continue
        cur = {
            "title": f"隐患排查依据-{small or mid or big}",
            "content": (f"检查项类别：{big} / {mid} / {small}。\n"
                        f"检查内容：{content}\n"
                        f"检查依据：{basis or '（见上）'}\n"
                        f"隐患级别：{level or '未分级'}"),
            "meta": {"big": big, "mid": mid, "small": small, "level": level or "未分级"},
        }
        docs.append(cur)
    wb.close()
    # 合并后续行后，content 里的换行保留
    return docs

def convert_worklist(path):
    """安环部工作清单: 行 → 文档（模块前向填充）"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    docs, last_mod = [], ""
    for r in ws.iter_rows(min_row=3, values_only=True):
        if not any(v is not None for v in r):
            continue
        vals = [clean(v) for v in r[:8]]
        seq, mod, mid, sub, task, freq, deadline = vals[0], vals[1], vals[2], vals[3], vals[4], vals[5], vals[6]
        if mod:
            last_mod = mod
        elif not task:
            continue
        mod_name = mod or last_mod
        docs.append({
            "title": f"安环部工作-{(mid or mod_name)[:12]}-{seq}",
            "content": (f"工作模块：{mod_name}{(' / ' + mid) if mid else ''}。\n"
                        f"工作任务：{task}\n"
                        f"工作频次：{freq or '按需'}\n"
                        f"完成时限：{deadline or '按制度要求'}"),
        })
    wb.close()
    return docs

def main():
    docs1 = convert_yiju(XLSX1)
    docs2 = convert_worklist(XLSX2)
    all_docs = docs1 + docs2
    for n, d in enumerate(all_docs, 1):
        d["title"] = f"{d['title']}#{n}"
    os.makedirs(os.path.dirname(SEED_JSON), exist_ok=True)
    json.dump(all_docs, open(SEED_JSON, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"转换完成: 依据大全 {len(docs1)} 条 + 工作清单 {len(docs2)} 条 = {len(all_docs)} 条 → {SEED_JSON}")

    if "--dry-run" in sys.argv:
        for d in all_docs[:3] + all_docs[-2:]:
            print("---", d["title"])
            print(d["content"][:150])
        return

    # 入库（BGE 语义嵌入）
    from mcp.knowledge_base import KnowledgeBase
    kb = KnowledgeBase(chroma_host="localhost", chroma_port=59999,
                       chroma_path=os.path.join(ROOT, "data", "chroma"))
    t0 = time.monotonic()
    n = kb.add_documents(all_docs)
    print(f"入库完成: {n} 片段, 耗时 {time.monotonic()-t0:.0f}s, 知识库总量 {kb.doc_count}")

if __name__ == "__main__":
    main()
